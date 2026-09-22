import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

root = Path(__file__).resolve().parents[1]
source = root / "requirements" / "nexa_requisitos.csv"
target = root / "requirements" / "nexa_requisitos.xlsx"

workbook = Workbook()
sheet = workbook.active
sheet.title = "Requisitos"
with source.open(encoding="utf-8", newline="") as file:
    for row in csv.reader(file):
        sheet.append(row)

for cell in sheet[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E78")
sheet.freeze_panes = "A2"
sheet.auto_filter.ref = sheet.dimensions
table = Table(displayName="RequisitosNexa", ref=sheet.dimensions)
table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
sheet.add_table(table)
for column in sheet.columns:
    width = min(max(len(str(cell.value or "")) for cell in column) + 2, 55)
    sheet.column_dimensions[column[0].column_letter].width = width
workbook.save(target)
print(target)